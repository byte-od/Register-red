import io
import re
from datetime import datetime
from typing import Dict, Any, List, Optional
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

try:
    from backend.models.schema import MonthData, ExpenseItem
except ImportError:
    from models.schema import MonthData, ExpenseItem

SUBCATEGORIES_CONFIG = {
    "food": {
        "id": "food",
        "name": "Alimentação & Delivery",
        "icon": "utensils",
        "color": "#EF4444",
        "bg_class": "bg-rose-500",
        "keywords": [
            "ifood", "99 food", "food", "lanche", "almoço", "almoco", "jantar",
            "restaurante", "pizza", "hamburguer", "mercado", "supermercado",
            "padaria", "confeitaria", "café", "cafe", "açaí", "acai", "pão", "churrasco",
            "marmita", "pastel", "sorvete", "bebida"
        ]
    },
    "transport": {
        "id": "transport",
        "name": "Transporte & Veículo",
        "icon": "car",
        "color": "#3B82F6",
        "bg_class": "bg-blue-500",
        "keywords": [
            "combustível", "combustivel", "gasolina", "etanol", "posto", "ipiranga",
            "shell", "óleo", "oleo", "carro", "veículo", "veiculo", "uber", "99 táxi",
            "99 taxi", "táxi", "taxi", "estacionamento", "pedágio", "pedagio", "lavagem",
            "pneu", "oficina", "mecanica", "mecânica", "multa", "ipva", "seguro auto"
        ]
    },
    "health": {
        "id": "health",
        "name": "Saúde, Farmácia & Pet",
        "icon": "heart-pulse",
        "color": "#10B981",
        "bg_class": "bg-emerald-500",
        "keywords": [
            "farmácia", "farmacia", "drogaria", "remédio", "remedio", "medicamento",
            "nicotina", "adesivo nicotina", "hospital", "médico", "medico", "consulta",
            "exame", "dentista", "clínica", "clinica", "hiperzoo", "pet", "veterinario",
            "veterinário", "ração", "racao", "psicólogo", "psicologo", "terapia"
        ]
    },
    "retail": {
        "id": "retail",
        "name": "Compras & Varejo",
        "icon": "shopping-bag",
        "color": "#F59E0B",
        "bg_class": "bg-amber-500",
        "keywords": [
            "lojinha", "loja", "marketplace", "diversos", "gráfica", "grafica",
            "shopping", "magazine", "amazon", "shopee", "mercado livre", "compra",
            "roupa", "eletronico", "eletrônico", "livro", "presente", "ferramenta",
            "papelaria", "utilidades", "vestuário", "vestuario"
        ]
    },
    "transfer": {
        "id": "transfer",
        "name": "Transferências & PIX",
        "icon": "send",
        "color": "#8B5CF6",
        "bg_class": "bg-purple-500",
        "keywords": [
            "transferência", "transferencia", "tatiane", "favorecido"
        ]
    },
    "other": {
        "id": "other",
        "name": "Outros Imprevistos",
        "icon": "zap",
        "color": "#64748B",
        "bg_class": "bg-slate-500",
        "keywords": []
    }
}

def classify_unforeseen_subcategory(name: str, notes: Optional[str] = None) -> Dict[str, str]:
    text = f"{name} {notes or ''}".lower().strip()
    
    # 1. Alimentação
    if any(k in text for k in SUBCATEGORIES_CONFIG["food"]["keywords"]):
        return SUBCATEGORIES_CONFIG["food"]

    # 2. Transporte
    if any(k in text for k in SUBCATEGORIES_CONFIG["transport"]["keywords"]):
        return SUBCATEGORIES_CONFIG["transport"]

    # 3. Saúde
    if any(k in text for k in SUBCATEGORIES_CONFIG["health"]["keywords"]):
        return SUBCATEGORIES_CONFIG["health"]

    # 4. Transferências
    if any(k in text for k in SUBCATEGORIES_CONFIG["transfer"]["keywords"]) or (
        "pix" in text and not any(k in text for k in ["marketplace", "loja", "mercado", "shop"])
    ):
        return SUBCATEGORIES_CONFIG["transfer"]

    # 5. Compras e Varejo
    if any(k in text for k in SUBCATEGORIES_CONFIG["retail"]["keywords"]):
        return SUBCATEGORIES_CONFIG["retail"]

    return SUBCATEGORIES_CONFIG["other"]

def extract_unforeseen_items(month: MonthData) -> List[Dict[str, Any]]:
    items: List[Dict[str, Any]] = []
    
    for cat in month.categories:
        is_gastos = cat.id.lower() == "gastos" or cat.name.lower() == "gastos"
        for it in cat.items:
            # Gastos são sempre imprevistos. Em outras categorias, se não for recorrente também é imprevisto.
            if is_gastos or not getattr(it, "is_recurring", True):
                sub = classify_unforeseen_subcategory(it.name, it.notes)
                items.append({
                    "id": it.id,
                    "name": it.name,
                    "amount": round(float(it.amount), 2),
                    "paid": it.paid,
                    "due_day": it.due_day or f"{month.year:04d}-{month.month_number:02d}-10",
                    "notes": it.notes,
                    "original_category_id": cat.id,
                    "original_category_name": cat.name,
                    "subcategory_id": sub["id"],
                    "subcategory_name": sub["name"],
                    "subcategory_color": sub["color"],
                    "subcategory_icon": sub["icon"],
                })

    # Ordena por valor decrescente por padrão
    items.sort(key=lambda x: x["amount"], reverse=True)
    return items

def build_unforeseen_report(month: MonthData, all_months: Optional[Dict[str, MonthData]] = None) -> Dict[str, Any]:
    items = extract_unforeseen_items(month)
    
    total_unforeseen = round(sum(it["amount"] for it in items), 2)
    total_expenses = round(float(month.summary.total_expenses), 2)
    total_income = round(float(month.summary.total_income), 2)
    current_balance = round(float(month.summary.balance), 2)
    balance_without_unforeseen = round(current_balance + total_unforeseen, 2)
    
    pct_expenses = round((total_unforeseen / total_expenses * 100), 1) if total_expenses > 0 else 0.0
    count = len(items)
    avg_ticket = round(total_unforeseen / count, 2) if count > 0 else 0.0

    highest_item = items[0] if items else None
    lowest_item = items[-1] if items else None

    # Subcategories aggregation
    subcats_map: Dict[str, Dict[str, Any]] = {}
    for cfg in SUBCATEGORIES_CONFIG.values():
        subcats_map[cfg["id"]] = {
            "id": cfg["id"],
            "name": cfg["name"],
            "icon": cfg["icon"],
            "color": cfg["color"],
            "bg_class": cfg["bg_class"],
            "total": 0.0,
            "percentage": 0.0,
            "count": 0,
            "items": []
        }

    for it in items:
        sid = it["subcategory_id"]
        if sid in subcats_map:
            subcats_map[sid]["total"] += it["amount"]
            subcats_map[sid]["count"] += 1
            subcats_map[sid]["items"].append(it)

    subcategories_list = []
    for sc in subcats_map.values():
        sc["total"] = round(sc["total"], 2)
        sc["percentage"] = round((sc["total"] / total_unforeseen * 100), 1) if total_unforeseen > 0 else 0.0
        if sc["count"] > 0:
            subcategories_list.append(sc)

    # Ordena subcategorias por maior volume
    subcategories_list.sort(key=lambda x: x["total"], reverse=True)

    # Daily distribution
    daily_sums: Dict[str, Dict[str, Any]] = {}
    for it in items:
        d_str = str(it["due_day"]).strip()
        match = re.search(r"(\d{4})-(\d{2})-(\d{2})", d_str)
        if match:
            day_num = int(match.group(3))
            full_date = d_str
        else:
            day_num = 10
            full_date = f"{month.year:04d}-{month.month_number:02d}-10"

        key = f"{day_num:02d}"
        if key not in daily_sums:
            daily_sums[key] = {
                "day": day_num,
                "date": full_date,
                "total": 0.0,
                "count": 0,
                "items": []
            }
        daily_sums[key]["total"] += it["amount"]
        daily_sums[key]["count"] += 1
        daily_sums[key]["items"].append(it)

    daily_distribution = sorted(daily_sums.values(), key=lambda x: x["day"])
    for d in daily_distribution:
        d["total"] = round(d["total"], 2)

    # Top 5 items
    top_items = items[:5]

    # Historical comparison with other months
    history_comparison: List[Dict[str, Any]] = []
    if all_months:
        sorted_m = sorted(all_months.values(), key=lambda m: (m.year, m.month_number))
        for m in sorted_m:
            m_items = extract_unforeseen_items(m)
            m_unforeseen = round(sum(it["amount"] for it in m_items), 2)
            m_expenses = round(float(m.summary.total_expenses), 2)
            m_pct = round((m_unforeseen / m_expenses * 100), 1) if m_expenses > 0 else 0.0
            history_comparison.append({
                "month_id": m.id,
                "month_title": m.title,
                "year": m.year,
                "month_number": m.month_number,
                "total_unforeseen": m_unforeseen,
                "total_expenses": m_expenses,
                "percentage_of_expenses": m_pct,
                "count": len(m_items),
                "is_current": m.id == month.id
            })

    # Smart Financial Insights
    insights: List[Dict[str, Any]] = []

    if total_unforeseen > 0:
        if pct_expenses > 25.0:
            insights.append({
                "type": "warning",
                "title": "Impacto orçamentário elevado",
                "message": f"Os gastos imprevistos somam {format_brl(total_unforeseen)} e respondem por {pct_expenses}% de todo o orçamento do mês. Especialistas recomendam manter imprevistos abaixo de 15%."
            })
        elif pct_expenses > 15.0:
            insights.append({
                "type": "info",
                "title": "Atenção aos gastos avulsos",
                "message": f"Os imprevistos representam {pct_expenses}% das despesas mensais ({format_brl(total_unforeseen)}). Uma margem moderada, mas com potencial de otimização."
            })
        else:
            insights.append({
                "type": "success",
                "title": "Excelente controle de imprevistos",
                "message": f"Seus gastos imprevistos representam apenas {pct_expenses}% das despesas totais ({format_brl(total_unforeseen)}), indicando ótima previsibilidade orçamentária."
            })

        # Saldo potencial sem imprevistos
        if balance_without_unforeseen > current_balance:
            dif = round(balance_without_unforeseen - current_balance, 2)
            insights.append({
                "type": "tip",
                "title": "Potencial de Poupança & Reserva",
                "message": f"Sem os imprevistos deste mês, seu saldo restante seria de {format_brl(balance_without_unforeseen)} (em vez de {format_brl(current_balance)}), liberando mais {format_brl(dif)} para economias ou investimentos."
            })

        # Subcategoria líder
        if subcategories_list:
            leader = subcategories_list[0]
            insights.append({
                "type": "info",
                "title": f"Maior concentração: {leader['name']}",
                "message": f"A área de '{leader['name']}' consumiu {format_brl(leader['total'])} ({leader['percentage']}% do total de imprevistos) distribuídos em {leader['count']} lançamentos."
            })

        # Detecção de alta frequência de compras pequenas (ex: lanches / deliveries)
        food_sub = next((s for s in subcategories_list if s["id"] == "food"), None)
        if food_sub and food_sub["count"] >= 3:
            insights.append({
                "type": "tip",
                "title": "Padrão recorrente de Alimentação/Delivery",
                "message": f"Identificamos {food_sub['count']} pedidos avulsos de alimentação/delivery ({format_brl(food_sub['total'])}). Considere fixar uma cota mensal de lazer/delivery nas despesas orçadas para evitar surpresas."
            })

        # Maior gasto único
        if highest_item and total_unforeseen > 0:
            pct_high = round((highest_item["amount"] / total_unforeseen * 100), 1)
            if pct_high >= 30.0:
                insights.append({
                    "type": "warning",
                    "title": f"Despesa pontual expressiva: '{highest_item['name']}'",
                    "message": f"O item '{highest_item['name']}' ({format_brl(highest_item['amount'])}) concentrou sozinho {pct_high}% de todos os imprevistos do mês."
                })
    else:
        insights.append({
            "type": "success",
            "title": "Nenhum gasto imprevisto registrado",
            "message": "Parabéns! Até o momento não há nenhum gasto não previsto registrado para este mês."
        })

    return {
        "month_id": month.id,
        "month_title": month.title,
        "year": month.year,
        "month_number": month.month_number,
        "total_unforeseen": total_unforeseen,
        "total_expenses": total_expenses,
        "percentage_of_expenses": pct_expenses,
        "total_income": total_income,
        "current_balance": current_balance,
        "balance_without_unforeseen": balance_without_unforeseen,
        "unforeseen_count": count,
        "average_ticket": avg_ticket,
        "highest_item": highest_item,
        "lowest_item": lowest_item,
        "subcategories": subcategories_list,
        "daily_distribution": daily_distribution,
        "top_items": top_items,
        "history_comparison": history_comparison,
        "insights": insights,
        "items": items
    }

def format_brl(val: float) -> str:
    return f"R$ {val:,.2f}".replace(",", "X").replace(".", ",").replace("X", ".")

def generate_unforeseen_excel(report: Dict[str, Any]) -> io.BytesIO:
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Resumo Imprevistos"

    # Styles
    title_font = Font(name="Arial", size=15, bold=True, color="1E293B")
    subtitle_font = Font(name="Arial", size=10, italic=True, color="64748B")
    section_font = Font(name="Arial", size=11, bold=True, color="FFFFFF")
    header_font = Font(name="Arial", size=10, bold=True, color="475569")
    bold_font = Font(name="Arial", size=10, bold=True, color="1E293B")
    regular_font = Font(name="Arial", size=10, color="1E293B")
    kpi_val_font = Font(name="Arial", size=13, bold=True, color="0F172A")
    alert_font = Font(name="Arial", size=10, bold=True, color="D97706")

    fill_dark = PatternFill(start_color="1E293B", end_color="1E293B", fill_type="solid")
    fill_header = PatternFill(start_color="F1F5F9", end_color="F1F5F9", fill_type="solid")
    fill_amber = PatternFill(start_color="FEF3C7", end_color="FEF3C7", fill_type="solid")
    fill_emerald = PatternFill(start_color="D1FAE5", end_color="D1FAE5", fill_type="solid")

    thin_border = Border(
        left=Side(style='thin', color='E2E8F0'),
        right=Side(style='thin', color='E2E8F0'),
        top=Side(style='thin', color='E2E8F0'),
        bottom=Side(style='thin', color='E2E8F0')
    )

    # 1. Title Banner
    ws.merge_cells("A1:F1")
    ws["A1"] = f"RELATÓRIO DE GASTOS IMPREVISTOS - {report['month_title'].upper()}"
    ws["A1"].font = title_font
    ws["A1"].alignment = Alignment(horizontal="center", vertical="center")
    ws.row_dimensions[1].height = 32

    ws.merge_cells("A2:F2")
    ws["A2"] = f"Análise detalhada de despesas avulsas e não planejadas | Gerado em {datetime.now().strftime('%d/%m/%Y às %H:%M')}"
    ws["A2"].font = subtitle_font
    ws["A2"].alignment = Alignment(horizontal="center", vertical="center")
    ws.row_dimensions[2].height = 20

    # 2. Executive KPI Summary Cards
    ws["A4"] = "Total Imprevistos"
    ws["B4"] = "% do Orçamento"
    ws["C4"] = "Qtd. Lançamentos"
    ws["D4"] = "Ticket Médio"
    ws["E4"] = "Saldo Atual"
    ws["F4"] = "Saldo s/ Imprevistos"

    for col in ["A", "B", "C", "D", "E", "F"]:
        ws[f"{col}4"].font = header_font
        ws[f"{col}4"].fill = fill_header
        ws[f"{col}4"].alignment = Alignment(horizontal="center")
        ws[f"{col}4"].border = thin_border

    ws["A5"] = report["total_unforeseen"]
    ws["B5"] = f"{report['percentage_of_expenses']}%"
    ws["C5"] = report["unforeseen_count"]
    ws["D5"] = report["average_ticket"]
    ws["E5"] = report["current_balance"]
    ws["F5"] = report["balance_without_unforeseen"]

    for col in ["A", "D", "E", "F"]:
        cell = ws[f"{col}5"]
        cell.font = kpi_val_font
        cell.number_format = 'R$ #,##0.00'
        cell.alignment = Alignment(horizontal="center")
        cell.border = thin_border

    ws["B5"].font = alert_font
    ws["B5"].alignment = Alignment(horizontal="center")
    ws["B5"].border = thin_border

    ws["C5"].font = kpi_val_font
    ws["C5"].alignment = Alignment(horizontal="center")
    ws["C5"].border = thin_border

    ws.row_dimensions[5].height = 26

    # 3. Subcategories Breakdown
    ws.merge_cells("A7:F7")
    ws["A7"] = "DISTRIBUIÇÃO POR SUB-NATUREZA TEMÁTICA"
    ws["A7"].font = section_font
    ws["A7"].fill = fill_dark
    ws.row_dimensions[7].height = 22

    ws["A8"] = "Subcategoria"
    ws["B8"] = "Total (R$)"
    ws["C8"] = "Participação (%)"
    ws["D8"] = "Qtd. Itens"
    ws.merge_cells("E8:F8")
    ws["E8"] = "Exemplos"

    for col in ["A", "B", "C", "D", "E", "F"]:
        ws[f"{col}8"].font = header_font
        ws[f"{col}8"].fill = fill_header
        ws[f"{col}8"].border = thin_border

    curr_row = 9
    for sc in report.get("subcategories", []):
        ws[f"A{curr_row}"] = sc["name"]
        ws[f"B{curr_row}"] = sc["total"]
        ws[f"B{curr_row}"].number_format = 'R$ #,##0.00'
        ws[f"C{curr_row}"] = f"{sc['percentage']}%"
        ws[f"D{curr_row}"] = sc["count"]
        
        sample_names = ", ".join(it["name"] for it in sc.get("items", [])[:3])
        ws.merge_cells(f"E{curr_row}:F{curr_row}")
        ws[f"E{curr_row}"] = sample_names

        for col in ["A", "B", "C", "D", "E", "F"]:
            ws[f"{col}{curr_row}"].font = regular_font
            ws[f"{col}{curr_row}"].border = thin_border

        curr_row += 1

    # 4. Full Items Listing
    curr_row += 1
    ws.merge_cells(f"A{curr_row}:F{curr_row}")
    ws[f"A{curr_row}"] = "LISTAGEM COMPLETA DOS GASTOS IMPREVISTOS"
    ws[f"A{curr_row}"].font = section_font
    ws[f"A{curr_row}"].fill = fill_dark
    ws.row_dimensions[curr_row].height = 22

    curr_row += 1
    headers = ["Data", "Descrição do Gasto", "Sub-natureza", "Categoria", "Valor (R$)", "Observação"]
    for idx, h in enumerate(headers, start=1):
        col_letter = get_column_letter(idx)
        ws[f"{col_letter}{curr_row}"] = h
        ws[f"{col_letter}{curr_row}"].font = header_font
        ws[f"{col_letter}{curr_row}"].fill = fill_header
        ws[f"{col_letter}{curr_row}"].border = thin_border
    curr_row += 1

    for it in report.get("items", []):
        ws[f"A{curr_row}"] = it.get("due_day") or ""
        ws[f"B{curr_row}"] = it.get("name") or ""
        ws[f"C{curr_row}"] = it.get("subcategory_name") or ""
        ws[f"D{curr_row}"] = it.get("original_category_name") or ""
        ws[f"E{curr_row}"] = it.get("amount") or 0.0
        ws[f"E{curr_row}"].number_format = 'R$ #,##0.00'
        ws[f"F{curr_row}"] = it.get("notes") or ""

        ws[f"A{curr_row}"].alignment = Alignment(horizontal="center")
        ws[f"E{curr_row}"].font = bold_font

        for col in ["A", "B", "C", "D", "E", "F"]:
            ws[f"{col}{curr_row}"].border = thin_border
        curr_row += 1

    # Adjust column widths
    ws.column_dimensions["A"].width = 15
    ws.column_dimensions["B"].width = 35
    ws.column_dimensions["C"].width = 25
    ws.column_dimensions["D"].width = 18
    ws.column_dimensions["E"].width = 18
    ws.column_dimensions["F"].width = 28

    buffer = io.BytesIO()
    wb.save(buffer)
    buffer.seek(0)
    return buffer
