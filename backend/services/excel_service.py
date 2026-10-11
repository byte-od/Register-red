import io
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter
from datetime import datetime

try:
    from backend.models.schema import MonthData
except ImportError:
    from models.schema import MonthData

def generate_excel_workbook(month: MonthData) -> io.BytesIO:
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = f"Orçamento {month.title}"

    # Styles
    title_font = Font(name="Arial", size=16, bold=True, color="1E293B")
    section_font = Font(name="Arial", size=12, bold=True, color="FFFFFF")
    header_font = Font(name="Arial", size=10, bold=True, color="475569")
    bold_font = Font(name="Arial", size=10, bold=True, color="1E293B")
    regular_font = Font(name="Arial", size=10, color="1E293B")
    success_font = Font(name="Arial", size=10, bold=True, color="16A34A")
    warning_font = Font(name="Arial", size=10, bold=True, color="DC2626")

    fill_dark = PatternFill(start_color="1E293B", end_color="1E293B", fill_type="solid")
    fill_header = PatternFill(start_color="F1F5F9", end_color="F1F5F9", fill_type="solid")
    fill_success = PatternFill(start_color="DCFCE7", end_color="DCFCE7", fill_type="solid")
    fill_warning = PatternFill(start_color="FEE2E2", end_color="FEE2E2", fill_type="solid")

    thin_border = Border(
        left=Side(style='thin', color='E2E8F0'),
        right=Side(style='thin', color='E2E8F0'),
        top=Side(style='thin', color='E2E8F0'),
        bottom=Side(style='thin', color='E2E8F0')
    )

    # Title Banner
    ws.merge_cells("A1:E1")
    ws["A1"] = f"GESTÃO FINANCEIRA - {month.title.upper()}"
    ws["A1"].font = title_font
    ws["A1"].alignment = Alignment(horizontal="center", vertical="center")
    ws.row_dimensions[1].height = 35

    # KPI Summary Cards
    ws["A3"] = "Receitas Totais"
    ws["B3"] = "Despesas Totais"
    ws["C3"] = "Total Pago"
    ws["D3"] = "A Pagar (Pendente)"
    ws["E3"] = "Saldo Restante"

    for col in ["A", "B", "C", "D", "E"]:
        ws[f"{col}3"].font = header_font
        ws[f"{col}3"].fill = fill_header
        ws[f"{col}3"].alignment = Alignment(horizontal="center")

    ws["A4"] = month.summary.total_income
    ws["B4"] = month.summary.total_expenses
    ws["C4"] = month.summary.total_paid
    ws["D4"] = month.summary.total_pending
    ws["E4"] = month.summary.balance

    for col in ["A", "B", "C", "D", "E"]:
        cell = ws[f"{col}4"]
        cell.font = bold_font
        cell.number_format = 'R$ #,##0.00'
        cell.alignment = Alignment(horizontal="center")
        cell.border = thin_border
    
    ws["C4"].font = success_font
    ws["D4"].font = warning_font
    ws["E4"].font = success_font if month.summary.balance >= 0 else warning_font

    current_row = 6

    # 1. Incomes Section
    ws.merge_cells(f"A{current_row}:C{current_row}")
    ws[f"A{current_row}"] = "RECEITAS & ENTRADAS"
    ws[f"A{current_row}"].font = section_font
    ws[f"A{current_row}"].fill = fill_dark
    ws.row_dimensions[current_row].height = 24
    current_row += 1

    ws[f"A{current_row}"] = "Descrição"
    ws[f"B{current_row}"] = "Valor (R$)"
    ws[f"C{current_row}"] = "Data / Vencimento"
    for col in ["A", "B", "C"]:
        ws[f"{col}{current_row}"].font = header_font
        ws[f"{col}{current_row}"].fill = fill_header
    current_row += 1

    for inc in month.incomes:
        ws[f"A{current_row}"] = inc.name
        ws[f"B{current_row}"] = inc.amount
        ws[f"B{current_row}"].number_format = 'R$ #,##0.00'
        ws[f"C{current_row}"] = inc.day or ""
        for col in ["A", "B", "C"]:
            ws[f"{col}{current_row}"].font = regular_font
            ws[f"{col}{current_row}"].border = thin_border
        current_row += 1

    # Income Total
    ws[f"A{current_row}"] = "Total Entradas"
    ws[f"A{current_row}"].font = bold_font
    ws[f"B{current_row}"] = month.summary.total_income
    ws[f"B{current_row}"].font = bold_font
    ws[f"B{current_row}"].number_format = 'R$ #,##0.00'
    ws[f"C{current_row}"] = ""
    for col in ["A", "B", "C"]:
        ws[f"{col}{current_row}"].fill = fill_header
        ws[f"{col}{current_row}"].border = thin_border
    current_row += 2

    # 2. Categories Section
    for cat in month.categories:
        ws.merge_cells(f"A{current_row}:E{current_row}")
        ws[f"A{current_row}"] = f"DESPESAS - {cat.name.upper()}"
        ws[f"A{current_row}"].font = section_font
        ws[f"A{current_row}"].fill = fill_dark
        ws.row_dimensions[current_row].height = 24
        current_row += 1

        headers = ["Status", "Descrição", "Valor (R$)", "Vencimento", "Observações"]
        for idx, h in enumerate(headers, start=1):
            col_letter = get_column_letter(idx)
            ws[f"{col_letter}{current_row}"] = h
            ws[f"{col_letter}{current_row}"].font = header_font
            ws[f"{col_letter}{current_row}"].fill = fill_header
        current_row += 1

        for it in cat.items:
            ws[f"A{current_row}"] = "PAGO" if it.paid else "PENDENTE"
            ws[f"A{current_row}"].fill = fill_success if it.paid else fill_warning
            ws[f"A{current_row}"].font = success_font if it.paid else warning_font
            ws[f"A{current_row}"].alignment = Alignment(horizontal="center")

            ws[f"B{current_row}"] = it.name
            ws[f"B{current_row}"].font = regular_font

            ws[f"C{current_row}"] = it.amount
            ws[f"C{current_row}"].font = bold_font
            ws[f"C{current_row}"].number_format = 'R$ #,##0.00'

            ws[f"D{current_row}"] = it.due_day or ""
            ws[f"D{current_row}"].font = regular_font
            ws[f"D{current_row}"].alignment = Alignment(horizontal="center")

            ws[f"E{current_row}"] = it.notes or ""
            ws[f"E{current_row}"].font = regular_font

            for col in ["A", "B", "C", "D", "E"]:
                ws[f"{col}{current_row}"].border = thin_border
            current_row += 1

        # Category Subtotal
        ws[f"A{current_row}"] = "Subtotal"
        ws[f"B{current_row}"] = f"Total {cat.name}"
        ws[f"B{current_row}"].font = bold_font
        ws[f"C{current_row}"] = cat.total
        ws[f"C{current_row}"].font = bold_font
        ws[f"C{current_row}"].number_format = 'R$ #,##0.00'
        ws[f"D{current_row}"] = f"Pago: R$ {cat.paid_total:,.2f}"
        ws[f"E{current_row}"] = f"Pendente: R$ {cat.pending_total:,.2f}"
        for col in ["A", "B", "C", "D", "E"]:
            ws[f"{col}{current_row}"].fill = fill_header
            ws[f"{col}{current_row}"].border = thin_border
        current_row += 2

    # Column widths
    ws.column_dimensions["A"].width = 16
    ws.column_dimensions["B"].width = 30
    ws.column_dimensions["C"].width = 20
    ws.column_dimensions["D"].width = 22
    ws.column_dimensions["E"].width = 25

    buffer = io.BytesIO()
    wb.save(buffer)
    buffer.seek(0)
    return buffer
